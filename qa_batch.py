"""Prueba de integración del análisis de varias fotos y su exportación."""

import gc
import os
import tempfile
import time
from types import SimpleNamespace
from unittest import mock

import cv2
import numpy as np
from openpyxl import load_workbook

import app


def make_image(path, centers):
    image = np.full((100, 200, 3), 235, dtype=np.uint8)
    for center in centers:
        cv2.circle(image, (center, 50), 10, (5, 5, 5), -1)
    assert cv2.imwrite(path, image)


def sheet_labels(sheet):
    """Pares etiqueta → valor (columnas A y B) de toda la hoja, sin depender
    de en qué fila quedó cada dato."""
    return {label: value for label, value in
            sheet.iter_rows(min_col=1, max_col=2, values_only=True) if label}


def pump(window, predicate, timeout=15):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        window.update()
        if predicate():
            return
        time.sleep(0.02)
    raise AssertionError("La interfaz no completó el análisis a tiempo")


def main():
    with tempfile.TemporaryDirectory(dir=os.path.dirname(__file__)) as directory:
        first = os.path.join(directory, "foto_1.png")
        second = os.path.join(directory, "foto_2.png")
        missing = os.path.join(directory, "no_existe.png")
        workbook = os.path.join(directory, "lote.xlsx")
        single_workbook = os.path.join(directory, "foto.xlsx")
        make_image(first, [50])
        make_image(second, [50, 140])

        partial = app._analyze_batch((first, missing), 40, 50, 2.0)
        assert list(partial["summaries"]) == [first]
        assert missing in partial["errors"]

        window = app.App()
        try:
            window.geometry("1300x790")
            window.update()
            window.um_per_pixel = 2.0
            window.show_drops()
            with mock.patch.object(app.filedialog, "askopenfilenames",
                                   return_value=(first, second, first)) as chooser:
                window.drops_zone._choose()
                chooser.assert_called_once()
            window.drops_zone._on_drop(
                SimpleNamespace(data=f"{{{first}}} {{{second}}}"))
            assert window.drops_paths == [first, second]
            assert window.drops_list.size() == 2

            with mock.patch.object(app.messagebox, "showerror"), \
                    mock.patch.object(app.messagebox, "showwarning"):
                window._submit_drops()
                pump(window, lambda: window._screen == "results" and
                     window._loading is None)
                assert [window.batch_results[path]["count"]
                        for path in (first, second)] == [1, 2]
                assert window.drops_path == first
                card = window._cards[0]["frame"]
                card_bottom = (card.winfo_rooty() + card.winfo_height() +
                               app.CARD_PAD + app.CARD_MARGIN)
                assert not window._scrollable
                assert card_bottom <= window.winfo_rooty() + window.winfo_height()

                window._select_result(1)
                pump(window, lambda: window.drops_path == second and
                     window._loading is None)
                assert window.result["count"] == 2
                assert "Foto 2 de 2" in window.result_nav_label.cget("text")

                window.focus_var.set(50)
                window._recalculate()
                pump(window, lambda: window._loading is None)
                assert window.batch_results[first]["count"] == 1
                assert window.batch_results[second]["count"] == 2

            window._write_batch_excel(workbook)
            book = load_workbook(workbook, data_only=True)
            try:
                assert book["Resumen"]["B2"].value == 2
                assert book["Resumen"]["B4"].value == 3
                labels = sheet_labels(book["Resumen"])
                assert labels["Oscuridad (%)"] == round(window.applied_darkness, 1)
                assert labels["Enfoque (%)"] == round(window.applied_focus, 1)
                assert book["Gotas"].max_row == 4
                assert book["Gotas"]["G2"].value == first
            finally:
                book.close()
            del book
            gc.collect()

            window.show_drops()
            window.drops_list.selection_set(0)
            window._remove_selected_drops()
            assert window.drops_paths == [second]
            window._submit_drops()
            pump(window, lambda: window._screen == "results" and
                 len(window.batch_results) == 1 and window._loading is None)
            assert window.result["count"] == 2
            assert not window.result_nav.winfo_manager()
            window._write_excel(single_workbook, window.result)
            book = load_workbook(single_workbook, data_only=True)
            try:
                assert book.sheetnames == ["Gotas", "Frecuencia"]
                assert book["Gotas"]["B3"].value == "foto_2.png"
                single_labels = sheet_labels(book["Gotas"])
                assert single_labels["Oscuridad (%)"] == round(window.applied_darkness, 1)
                assert single_labels["Enfoque (%)"] == round(window.applied_focus, 1)
            finally:
                book.close()
            del book
            gc.collect()
            window.show_drops()
            window._clear_drops_file()
            assert not window.drops_paths
        finally:
            window.destroy()
    print("PASS: selección múltiple, análisis, navegación, enfoque y Excel")


if __name__ == "__main__":
    main()
