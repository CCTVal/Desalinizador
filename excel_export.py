"""Exportación a Excel de la foto visible y del lote completo.

Las funciones reciben datos ya calculados (sin Tk), para poder probarlas y
para que la interfaz solo se ocupe de reunir el estado."""

import os

import numpy as np


def batch_totals(batch_results):
    """Totales globales del lote: se muestran en pantalla y en el Excel."""
    widths = [width for image in batch_results.values()
              for width in image["widths"]]
    return {
        "photos": len(batch_results),
        "count": sum(image["count"] for image in batch_results.values()),
        "mean": float(np.mean(widths)) if widths else None,
        "std": float(np.std(widths)) if widths else None,
        "widths": widths,
    }


def write_single(path, res, info):
    """Una sola foto. `info` trae: title, drops_name, calib_name,
    um_per_pixel, darkness, focus."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    unit = res["unit"]
    upp = info["um_per_pixel"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Gotas"

    head_fill = PatternFill("solid", fgColor="2E31C6")   # índigo CCTVal
    head_font = Font(bold=True, color="FFFFFF")
    bold = Font(bold=True)

    ws["A1"] = info["title"]
    ws["A1"].font = Font(bold=True, size=14, color="2E31C6")

    summary = [
        ("Archivo de gotas", info["drops_name"] or "—"),
        ("Imagen de referencia", info["calib_name"] or "—"),
        ("Escala", f"{upp:.4f} µm/píxel" if upp
         else "Sin calibración (píxeles)"),
        ("Oscuridad (%)", round(info["darkness"], 1)),
        ("Enfoque (%)", round(info["focus"], 1)),
        ("Unidad de medida", unit),
        ("Gotas reconocidas", res["count"]),
        (f"Tamaño promedio ({unit})", round(res["mean_width"], 2)),
        (f"Desviación estándar ({unit})", round(res["std_width"], 2)),
    ]
    r = 3
    for key, value in summary:
        ws.cell(row=r, column=1, value=key).font = bold
        ws.cell(row=r, column=2, value=value)
        r += 1

    r += 1
    headers = ["#", f"Ancho ({unit})", "Ancho (px)", "Alto (px)",
               "Área (px²)"]
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=r, column=col, value=text)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center")
    header_row = r
    r += 1
    for i, c in enumerate(res.get("circles", []), start=1):
        wpx = c["max_width"]
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2,
                value=round(wpx * upp, 2) if upp else wpx)
        ws.cell(row=r, column=3, value=wpx)
        ws.cell(row=r, column=4, value=c["max_height"])
        ws.cell(row=r, column=5, value=c["pixel_count"])
        r += 1

    for col, width in enumerate((6, 16, 12, 12, 14), start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    # Segunda hoja: distribución de frecuencias (mismos 30 bins que el
    # histograma de la pantalla de resultados).
    fs = wb.create_sheet("Frecuencia")
    f_headers = [f"Ancho desde ({unit})", f"Ancho hasta ({unit})",
                 "Frecuencia"]
    for col, text in enumerate(f_headers, start=1):
        cell = fs.cell(row=1, column=col, value=text)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center")
    widths_data = res.get("widths", [])
    if widths_data:
        counts, edges = np.histogram(widths_data, bins=30)
        for i, n in enumerate(counts):
            fs.cell(row=i + 2, column=1, value=round(float(edges[i]), 2))
            fs.cell(row=i + 2, column=2, value=round(float(edges[i + 1]), 2))
            fs.cell(row=i + 2, column=3, value=int(n))
    for col, width in enumerate((18, 18, 12), start=1):
        fs.column_dimensions[get_column_letter(col)].width = width
    fs.freeze_panes = fs.cell(row=2, column=1)

    wb.save(path)


def write_batch(path, drops_paths, batch_results, batch_errors, info):
    """Lote de fotos. `info` trae: title, unit, calib_path, um_per_pixel,
    darkness, focus."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    unit = info["unit"]
    totals = batch_totals(batch_results)
    wb = Workbook()
    summary = wb.active
    summary.title = "Resumen"
    drops = wb.create_sheet("Gotas")
    frequency = wb.create_sheet("Frecuencia")

    summary["A1"] = info["title"] + " · lote de fotos"
    summary["A1"].font = Font(bold=True, size=14, color="2E31C6")
    upp = info["um_per_pixel"]
    metadata = [
        ("Fotos seleccionadas", len(drops_paths)),
        ("Fotos analizadas", totals["photos"]),
        ("Gotas totales", totals["count"]),
        (f"Promedio global ({unit})", round(totals["mean"], 2)
         if totals["mean"] is not None else 0),
        (f"Desviación global ({unit})", round(totals["std"], 2)
         if totals["std"] is not None else 0),
        ("Escala", f"{upp:.4f} µm/píxel" if upp
         else "Sin calibración (píxeles)"),
        ("Oscuridad (%)", round(info["darkness"], 1)),
        ("Enfoque (%)", round(info["focus"], 1)),
        ("Imagen de referencia", info["calib_path"] or "—"),
    ]
    for row, (label, value) in enumerate(metadata, start=2):
        summary.cell(row, 1, label).font = Font(bold=True)
        summary.cell(row, 2, value)
    table_row = len(metadata) + 3

    fill = PatternFill("solid", fgColor="2E31C6")
    white = Font(bold=True, color="FFFFFF")

    def header(sheet, row, labels):
        for col, label in enumerate(labels, start=1):
            cell = sheet.cell(row, col, label)
            cell.fill = fill
            cell.font = white

    header(summary, table_row, ["Archivo", "Estado", "Gotas",
                                f"Promedio ({unit})", f"Desviación ({unit})",
                                "Ruta", "Error"])
    header(drops, 1, ["Archivo", "#", f"Ancho ({unit})",
                      "Ancho (px)", "Alto (px)", "Área (px²)", "Ruta"])
    drop_row = 2
    for row, image_path in enumerate(drops_paths, start=table_row + 1):
        res = batch_results.get(image_path)
        summary.cell(row, 1, os.path.basename(image_path))
        summary.cell(row, 2, "Analizada" if res else "Error")
        summary.cell(row, 6, image_path)
        if not res:
            summary.cell(row, 7, batch_errors.get(image_path,
                                                  "No se pudo analizar"))
            continue
        summary.cell(row, 3, res["count"])
        summary.cell(row, 4, round(res["mean_width"], 2))
        summary.cell(row, 5, round(res["std_width"], 2))
        for index, circle in enumerate(res.get("circles", []), start=1):
            drops.cell(drop_row, 1, os.path.basename(image_path))
            drops.cell(drop_row, 2, index)
            drops.cell(drop_row, 3, round(res["widths"][index - 1], 2))
            drops.cell(drop_row, 4, circle["max_width"])
            drops.cell(drop_row, 5, circle["max_height"])
            drops.cell(drop_row, 6, circle["pixel_count"])
            drops.cell(drop_row, 7, image_path)
            drop_row += 1

    header(frequency, 1, [f"Ancho desde ({unit})",
                          f"Ancho hasta ({unit})", "Frecuencia"])
    if totals["widths"]:
        counts, edges = np.histogram(totals["widths"], bins=30)
        for row, count in enumerate(counts, start=2):
            frequency.cell(row, 1, round(float(edges[row - 2]), 2))
            frequency.cell(row, 2, round(float(edges[row - 1]), 2))
            frequency.cell(row, 3, int(count))

    for sheet in (summary, drops):
        sheet.column_dimensions["A"].width = 32
        sheet.column_dimensions["B"].width = 20
        sheet.column_dimensions["C"].width = 16
        sheet.column_dimensions["D"].width = 20
        sheet.column_dimensions["E"].width = 20
        sheet.column_dimensions["F"].width = 45
        sheet.column_dimensions["G"].width = 45
    frequency.column_dimensions["A"].width = 20
    frequency.column_dimensions["B"].width = 20
    frequency.column_dimensions["C"].width = 14
    summary.freeze_panes = f"A{table_row + 1}"
    drops.freeze_panes = "A2"
    frequency.freeze_panes = "A2"
    wb.save(path)
